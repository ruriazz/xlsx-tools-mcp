import pytest
from xlsx_tools_mcp.io.writer import create_workbook, append_rows, create_table, create_chart
import openpyxl

def test_create_table_chart(tmp_path):
    path = str(tmp_path / "test_out.xlsx")
    create_workbook(path, sheets=["Sheet1"])
    append_rows(path, "Sheet1", [
        ["Month", "Sales"],
        ["Jan", 100],
        ["Feb", 150],
        ["Mar", 200]
    ])
    
    # Test create table
    result = create_table(path, "Sheet1", "A1:B4", "SalesTable")
    assert result["saved"] is True
    
    wb = openpyxl.load_workbook(path)
    assert "SalesTable" in wb["Sheet1"].tables
    
    # Test create chart
    res_chart = create_chart(path, "Sheet1", "bar", data_range="B1:B4", categories_range="A2:A4", title="Sales Chart")
    assert res_chart["saved"] is True
    
    wb2 = openpyxl.load_workbook(path)
    assert len(wb2["Sheet1"]._charts) == 1
    chart = wb2["Sheet1"]._charts[0]
    assert chart.title.tx.rich.p[0].r[0].t == "Sales Chart"
